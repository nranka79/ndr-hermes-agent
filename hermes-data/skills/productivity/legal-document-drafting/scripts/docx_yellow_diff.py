#!/usr/bin/env python3
"""Rewrite a .docx from a replacement spec, highlighting ONLY the changed spans in yellow.

Built for the DRAAS "re-issue this sale-deed draft for a new plot / new buyer, show me the
updates in yellow" workflow (see references/plotted-layout-sale-deed-update.md).

Why difflib instead of "highlight the whole paragraph": the user asked to see *what changed*.
SequenceMatcher opcodes give exactly the changed spans, so unchanged boilerplate in the same
paragraph stays white and review is fast.

Usage
-----
    # apply the spec and write a new docx
    python3 docx_yellow_diff.py src.docx out.docx spec.json

    # verify an already-written docx (highlight counts + stale-string scan)
    python3 docx_yellow_diff.py --verify-only out.docx spec.json

spec.json
---------
    {
      "exact":  { "<exact old paragraph text>": "<full new paragraph text>" },
      "global": { "old substring": "new substring" },
      "stale":  ["MANJUNATH MANOHAR SINGH", "21,77,200", "Plot No. 34"]
    }

`exact` is matched against the paragraph's `.text` (stripped) - use it for whole-block
rewrites. `global` substring swaps are then applied to every other paragraph. `stale` is the
list of strings that must NOT survive anywhere afterwards - it is the regression check that
stops a half-updated amount or an old party name shipping to a Sub-Registrar.

Dependencies
------------
python-docx. In this environment:
    uv venv /tmp/dv --system-site-packages
    uv pip install --python /tmp/dv/bin/python python-docx
then run with /tmp/dv/bin/python.

Pitfalls this script already handles (do not "simplify" them away)
-----------------------------------------------------------------
* `doc.paragraphs` EXCLUDES paragraphs inside tables. Signature/execution blocks and page
  footers are usually tables, so the old vendee's name would survive in the footer of every
  page. This walks `body.iter(w:p)` and each section's header/footer element tree instead.
* `<w:highlight w:val="yellow"/>` must be inserted in OOXML rPr schema order (after
  w:sz/w:szCs, before w:rtl). Appending it to the end of rPr can render nothing. `set_highlight`
  does an ordered insert.
* Use w:highlight, never w:shd - w:shd is often invisible in Google Docs.
* Run-level rebuild drops any hyperlink runs in the paragraph. For a deed body that is fine;
  do not point this at a document whose formatting you cannot re-check.
"""

import argparse
import copy
import difflib
import json
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

HL = "w:highlight"

# OOXML CT_RPr child order (the subset that matters for run properties).
RPR_ORDER = [
    "w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps", "w:smallCaps",
    "w:strike", "w:dstrike", "w:outline", "w:shadow", "w:emboss", "w:imprint",
    "w:noProof", "w:snapToGrid", "w:vanish", "w:webHidden", "w:color", "w:spacing",
    "w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect",
    "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
    "w:eastAsianLayout", "w:specVanish", "w:oMath",
]


def iter_paragraphs(doc):
    """Every paragraph in the document, INCLUDING table cells, headers and footers."""
    for el in doc.element.body.iter(qn("w:p")):
        yield Paragraph(el, doc)
    for sec in doc.sections:
        for part in (sec.header, sec.footer):
            if part is None:
                continue
            for el in part._element.iter(qn("w:p")):
                yield Paragraph(el, part)


def set_highlight(rpr):
    """Force a yellow highlight on an rPr, in schema-valid position."""
    for h in rpr.findall(qn(HL)):
        rpr.remove(h)
    h = OxmlElement(HL)
    h.set(qn("w:val"), "yellow")
    idx = RPR_ORDER.index(HL)
    for child in rpr:
        name = "w:" + child.tag.split("}")[-1]
        if name in RPR_ORDER and RPR_ORDER.index(name) > idx:
            child.addprevious(h)
            return
    rpr.append(h)


def rebuild(p, new_text):
    """Replace the paragraph's text, yellow-highlighting only the changed spans."""
    runs = p.runs
    if not runs:
        return False
    old = p.text
    if old == new_text:
        return False
    base_rpr = runs[0]._element.find(qn("w:rPr"))
    base_rpr = copy.deepcopy(base_rpr) if base_rpr is not None else None
    for r in list(runs):
        r._element.getparent().remove(r._element)
    sm = difflib.SequenceMatcher(None, old, new_text, autojunk=False)
    for tag, _i1, _i2, j1, j2 in sm.get_opcodes():
        seg = new_text[j1:j2]
        if not seg:
            continue
        nr = OxmlElement("w:r")
        if base_rpr is not None:
            nr.append(copy.deepcopy(base_rpr))
        if tag != "equal":
            rpr = nr.find(qn("w:rPr"))
            if rpr is None:
                rpr = OxmlElement("w:rPr")
                nr.insert(0, rpr)
            set_highlight(rpr)
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = seg
        nr.append(t)
        p._p.append(nr)
    return True


def apply_spec(doc, spec):
    exact = {k.strip(): v for k, v in spec.get("exact", {}).items()}
    glob = spec.get("global", {})
    changed = 0
    for p in iter_paragraphs(doc):
        t = p.text
        if not t.strip():
            continue
        key = t.strip()
        if key in exact:
            new = exact[key]
        else:
            new = t
            for a, b in glob.items():
                new = new.replace(a, b)
        if new != t and rebuild(p, new):
            changed += 1
    return changed


def report(doc, stale=()):
    highlighted = total = 0
    hits = []
    for p in iter_paragraphs(doc):
        for r in p.runs:
            total += 1
            rpr = r._element.find(qn("w:rPr"))
            if rpr is not None and rpr.find(qn(HL)) is not None:
                highlighted += 1
        text = p.text
        for s in stale:
            if s.lower() in text.lower():
                hits.append((s, text.strip()[:140]))
    print("runs=%d highlighted=%d" % (total, highlighted))
    if not highlighted:
        print("WARNING: nothing is highlighted - check the spec matched.")
    for s, text in hits:
        print("  STALE[%s] %s" % (s, text))
    if stale and not hits:
        print("stale scan: clean")
    return highlighted, hits


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("src")
    ap.add_argument("out", nargs="?")
    ap.add_argument("spec")
    ap.add_argument("--verify-only", action="store_true")
    args = ap.parse_args()

    spec = json.load(open(args.spec, encoding="utf-8"))
    stale = spec.get("stale", [])

    doc = Document(args.src)
    if not args.verify_only:
        if not args.out:
            ap.error("out is required unless --verify-only")
        changed = apply_spec(doc, spec)
        doc.save(args.out)
        print("paragraphs changed:", changed)
        print("wrote", args.out)
        doc = Document(args.out)

    _hl, hits = report(doc, stale)
    return 1 if hits else 0


if __name__ == "__main__":
    sys.exit(main())
