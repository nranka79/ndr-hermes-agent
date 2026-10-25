# -*- coding: utf-8 -*-
"""
Starter builder: one content list -> matching .docx + .pdf (+ .txt) for a legal draft.

Copy this file, replace CONTENT with the real clause text, and run it twice:

    /opt/data/.venv-docx/bin/python    build.py     # writes .docx and .txt
    /opt/data/.venv-render/bin/python -c "import build; print(build.build_pdf())"

Why two interpreters: python-docx lives in /opt/data/.venv-docx and weasyprint in
/opt/data/.venv-render. There is no LibreOffice/pandoc on this host, so HTML -> PDF via
weasyprint is the PDF path. NEVER deliver the intermediate .html to the user.

Block kinds: title, sub, center, centerb, h1, h2, p, b, small, gap, pagebreak, rule,
             table (rows, optional widths), sigblock.
A block starting with '\t' in its text is rendered as an indented sub-paragraph.
"""
import os

OUT = os.environ.get('OUT_BASENAME', '/tmp/legal_draft')
OUT_DOCX, OUT_PDF, OUT_TXT = OUT + '.docx', OUT + '.pdf', OUT + '.txt'

# ----------------------------------------------------------------- content
CONTENT = [
    ('title', 'TITLE OF THE DEED'),
    ('sub', 'PROJECT — VILLAGE, TALUK, DISTRICT, STATE'),
    ('p', 'This DEED is made and executed at ________ on this the ____ day of ________.'),
    ('centerb', 'BETWEEN'),
    ('p', '1.\t<party 1 full identification: name, s/o, age, DOB, Aadhaar, PAN, address>'),
    ('p', '\t(Hereinafter referred to as the "<SHORT NAME>" ...)'),
    ('center', 'OF THE FIRST PART'),
    ('gap', ''),
    ('centerb', 'AND'),
    ('p', '2.\t<party 2 full identification>'),
    ('center', 'OF THE SECOND PART'),
    ('pagebreak', ''),
    ('h1', 'RECITALS'),
    ('p', 'WHEREAS ...'),
    ('p', 'NOW, THEREFORE, IN CONSIDERATION OF THE FOREGOING RECITALS AND THE MUTUAL '
          'COVENANTS HEREIN CONTAINED, THE PARTIES HEREBY AGREE AS FOLLOWS:'),
    ('h1', '1.\tDEFINITIONS AND INTERPRETATION'),
    ('p', '1.1\t...'),
    ('h1', 'SIGNATURES'),
    ('sigblock', ''),
]

SIGBLOCK = [
    'FIRST PART',
    '',
    '<name>',
    '______________________________________',
    '(Signature)',
    '',
    'SECOND PART',
    '',
    'For <entity>',
    '______________________________________',
    '<name>, Director',
    '',
    'WITNESSES',
    '1. ______________________________   2. ______________________________',
    '',
    'DRAFTED BY: Self',
]

# ------------------------------------------------------------ plain text
def build_txt():
    out = []
    for blk in CONTENT:
        k = blk[0]
        if k == 'table':
            for r in blk[1][0]:
                out.append(' | '.join(str(c) for c in r))
            out.append('')
        elif k == 'sigblock':
            out.extend(SIGBLOCK); out.append('')
        elif k == 'rule':
            out.append('_' * 78)
        elif k == 'pagebreak':
            out.append('\f')
        elif k == 'gap':
            out.append('')
        else:
            out.append(blk[1])
    open(OUT_TXT, 'w').write('\n'.join(out))
    return OUT_TXT

# ------------------------------------------------------------------- docx
def build_docx():
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2.2)

    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.15

    def para(text, size=11, bold=False, align='just', indent=0, sb=0, sa=6, italic=False):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(sb), Pt(sa), 1.15
        if indent:
            pf.left_indent = Cm(indent)
        pf.alignment = {'just': WD_ALIGN_PARAGRAPH.JUSTIFY, 'center': WD_ALIGN_PARAGRAPH.CENTER,
                        'left': WD_ALIGN_PARAGRAPH.LEFT}[align]
        r = p.add_run(text)
        r.font.size, r.bold, r.italic, r.font.name = Pt(size), bold, italic, 'Times New Roman'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        return p

    def add_table(rows, widths=None):
        tb = doc.add_table(rows=0, cols=len(rows[0]))
        tb.style = 'Table Grid'
        for ri, row in enumerate(rows):
            cells = tb.add_row().cells
            for ci, val in enumerate(row):
                p = cells[ci].paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                r = p.add_run(str(val))
                r.font.size, r.font.name, r.bold = Pt(9.5), 'Times New Roman', (ri == 0)
        if widths:                                   # widths in quarter-centimetres
            for ri in range(len(tb.rows)):
                for ci, w in enumerate(widths):
                    tb.rows[ri].cells[ci].width = Cm(w / 4.0)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    for blk in CONTENT:
        k, txt = blk[0], blk[1] if len(blk) > 1 else ''
        if k == 'title':    para(txt, size=16, bold=True, align='center', sa=8)
        elif k == 'sub':    para(txt, size=10, align='center', sa=4, italic=True)
        elif k == 'center': para(txt, bold=True, align='center', sb=6, sa=6)
        elif k == 'centerb':para(txt, size=12, bold=True, align='center', sb=8, sa=8)
        elif k == 'h1':     para(txt, size=12, bold=True, align='left', sb=14, sa=8)
        elif k == 'h2':     para(txt, size=11, bold=True, align='left', sb=10, sa=6)
        elif k == 'p':
            para(txt.lstrip('\t'), indent=0.9 if txt.startswith('\t') else 0,
                 sa=4 if txt.startswith('\t') else 6)
        elif k == 'b':      para(txt, bold=True, align='left', sb=8, sa=6)
        elif k == 'small':  para(txt, size=9.5, sa=4)
        elif k == 'rule':   para('_' * 78, align='center', sa=6)
        elif k == 'gap':    doc.add_paragraph().paragraph_format.space_after = Pt(4)
        elif k == 'pagebreak': doc.add_page_break()
        elif k == 'table':
            rows, widths = blk[1] if len(blk[1]) == 2 else (blk[1][0], None)
            add_table(rows, widths)
        elif k == 'sigblock':
            for line in SIGBLOCK:
                if line == '':
                    doc.add_paragraph().paragraph_format.space_after = Pt(2)
                else:
                    para(line, align='left', sa=2,
                         bold=(line.isupper() and len(line) < 60))
    doc.save(OUT_DOCX)
    return OUT_DOCX

# -------------------------------------------------------------------- pdf
CSS = """
@page { size: A4; margin: 20mm 18mm 20mm 18mm;
        @bottom-center { content: "Page " counter(page) " of " counter(pages);
                         font-family: 'Times New Roman', serif; font-size: 8pt; color:#444; } }
body { font-family: 'Times New Roman', Times, serif; font-size: 10.5pt;
       line-height: 1.35; text-align: justify; }
h1.title { font-size:16pt; text-align:center; margin:0 0 6pt 0; }
p.sub { font-size:9.5pt; text-align:center; font-style:italic; margin:0 0 4pt 0; }
p.center { text-align:center; font-weight:bold; margin:8pt 0; }
p.centerb { text-align:center; font-weight:bold; font-size:11.5pt; margin:10pt 0; }
h2.h1 { font-size:12pt; margin:14pt 0 7pt 0; page-break-after:avoid; text-align:left; }
h3.h2 { font-size:10.5pt; margin:10pt 0 5pt 0; text-align:left; }
p { margin:0 0 5pt 0; }
p.ind { margin-left:22pt; }
p.bold { font-weight:bold; text-align:left; margin:8pt 0 5pt 0; }
p.small { font-size:9pt; }
p.rule { text-align:center; letter-spacing:2px; margin:6pt 0; }
hr.pb { border:none; page-break-after:always; margin:0; }
table { border-collapse:collapse; width:100%; margin:6pt 0 10pt 0; font-size:8.8pt; }
th, td { border:0.6pt solid #333; padding:3pt 4pt; vertical-align:top; text-align:left; }
th { background:#eeeeee; font-weight:bold; }
p.sig { margin:0 0 2pt 0; text-align:left; }
div.gap { height:6pt; }
"""

def build_pdf():
    import html as _html
    from weasyprint import HTML
    esc = lambda x: _html.escape(str(x))

    def render_table(rows):
        o = ['<table>']
        for ri, row in enumerate(rows):
            o.append('<tr>')
            for c in row:
                tag = 'th' if ri == 0 else 'td'
                o.append('<%s>%s</%s>' % (tag, esc(c), tag))
            o.append('</tr>')
        o.append('</table>')
        return ''.join(o)

    parts = ['<html><head><meta charset="utf-8"><style>%s</style></head><body>' % CSS]
    for blk in CONTENT:
        k, txt = blk[0], blk[1] if len(blk) > 1 else ''
        if k == 'title':    parts.append('<h1 class="title">%s</h1>' % esc(txt))
        elif k == 'sub':    parts.append('<p class="sub">%s</p>' % esc(txt))
        elif k == 'center': parts.append('<p class="center">%s</p>' % esc(txt))
        elif k == 'centerb':parts.append('<p class="centerb">%s</p>' % esc(txt))
        elif k == 'h1':     parts.append('<h2 class="h1">%s</h2>' % esc(txt))
        elif k == 'h2':     parts.append('<h3 class="h2">%s</h3>' % esc(txt))
        elif k == 'p':
            cls = 'ind' if txt.startswith('\t') else ''
            parts.append('<p class="%s">%s</p>' % (cls, esc(txt.lstrip('\t'))))
        elif k == 'b':      parts.append('<p class="bold">%s</p>' % esc(txt))
        elif k == 'small':  parts.append('<p class="small">%s</p>' % esc(txt))
        elif k == 'rule':   parts.append('<p class="rule">%s</p>' % esc('_' * 78))
        elif k == 'gap':    parts.append('<div class="gap"></div>')
        elif k == 'pagebreak': parts.append('<hr class="pb">')
        elif k == 'table':
            rows = blk[1][0] if isinstance(blk[1], tuple) else blk[1]
            parts.append(render_table(rows))
        elif k == 'sigblock':
            for line in SIGBLOCK:
                parts.append('<p class="sig">%s</p>' % (esc(line) if line else '&nbsp;'))
    parts.append('</body></html>')
    HTML(string=''.join(parts)).write_pdf(OUT_PDF)
    return OUT_PDF

if __name__ == '__main__':
    print(build_txt())
    print(build_docx())
    for f in (OUT_DOCX, OUT_TXT):
        print(f, os.path.getsize(f))
